from __future__ import annotations

import argparse
import json
import os
import re
import urllib.error
import urllib.request
from datetime import datetime, timezone
from http.server import SimpleHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook

ROOT = Path(__file__).resolve().parent
ENV_PATH = ROOT / ".env"
FAQ_PATH = ROOT / "ulearning_teacher_faq.json"
MANUAL_2026_FAQ_PATH = ROOT / "manual_2026_faq.json"
PROMPT_PATH = ROOT / "deepseek_prompt.md"
HELP_KB_PATH = ROOT / "help知识库.docx"
CHUNK_KB_PATH = ROOT / "切片读取知识库" / "knowledge_chunks.jsonl"
SCREENSHOT_KB_PATH = ROOT / "切片读取知识库" / "screenshot_tutorial_kb.jsonl"
SCREENSHOT_EN_TRANSLATIONS_PATH = ROOT / "screenshot_guide_translations_en.json"
MANUAL_2026_GUIDES_PATH = ROOT / "manual_2026_guides.json"
PRECONDITION_PATH = ROOT / "不完全前置条件功能.txt"
UNKNOWN_LOG_PATH = ROOT / "unanswered_questions.jsonl"
FEEDBACK_DIR = ROOT / "feedback"
SCREENSHOT_DIR = ROOT / "截图教程"
VIDEO_DIRS = (ROOT / "视频教程", ROOT / "指导视频")
IMAGE_EXTENSIONS = {".png", ".jpg", ".jpeg", ".webp", ".bmp"}
VIDEO_EXTENSIONS = {".mp4", ".webm", ".mov", ".m4v"}
HELPFUL_FEEDBACK_TABLE_PATH = FEEDBACK_DIR / "有帮助反馈表.xlsx"
UNHELPFUL_FEEDBACK_TABLE_PATH = FEEDBACK_DIR / "无帮助反馈表.xlsx"
HANDOFF_FEEDBACK_TABLE_PATH = FEEDBACK_DIR / "转人工反馈表.xlsx"
FEEDBACK_FIELDS = ["用户", "时间", "问题", "记录ID"]
VISIBLE_FEEDBACK_FIELDS = ["用户", "时间", "问题"]
_HELP_KB_CACHE: list[dict[str, Any]] = []
_HELP_KB_MTIME_NS: int | None = None
_JSONL_KB_CACHE: dict[str, tuple[int, list[dict[str, Any]]]] = {}
_PRECONDITION_CACHE: list[dict[str, Any]] = []
_PRECONDITION_MTIME_NS: int | None = None

BUILTIN_GUIDE_ITEMS = [
    {
        "id": "course-ai-assistant",
        "category": "course",
        "question": "如何使用和设置课程主页 AI 助手？",
        "summary": "课程工作台提供教师备课助手和学生 AI 助教入口。",
        "keywords": [
            "AI助手",
            "AI助教",
            "学生AI助教",
            "教师备课助手",
            "备课助手",
            "建立自己的AI助手",
            "创建AI助手",
            "课程主页AI",
            "课程工作台",
        ],
        "answer": {
            "entry": "进入目标课程，在左侧打开【工作台】。",
            "steps": [
                "在工作台顶部选择【教师备课助手】或【学生AI助教】",
                "需要提问时，在输入框填写要求，也可使用附件或语音入口",
                "需要新建对话或添加助手时，查看右上角的【+】按钮",
                "需要修改助手配置时，打开右上角的齿轮设置",
            ],
            "check": "页面显示对应助手的对话区和功能入口。",
            "risk": "如果用户说“建立自己的AI助手”，先确认是配置课程内已有助手，还是要开发独立于 uLearning 的新助手；前者属于 uLearning 教师操作。",
        },
    }
]


def load_env_file(path: Path) -> None:
    if not path.exists():
        return
    for raw_line in path.read_text(encoding="utf-8").splitlines():
        line = raw_line.strip()
        if not line or line.startswith("#") or "=" not in line:
            continue
        key, value = line.split("=", 1)
        os.environ.setdefault(key.strip(), value.strip().strip('"').strip("'"))


def load_faq_items() -> list[dict[str, Any]]:
    try:
        data = json.loads(FAQ_PATH.read_text(encoding="utf-8"))
    except OSError:
        return []
    if not isinstance(data, list):
        return []
    items = [item for item in data if isinstance(item, dict)]
    try:
        manual_data = json.loads(MANUAL_2026_FAQ_PATH.read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError):
        manual_data = []
    if isinstance(manual_data, list):
        items.extend(item for item in manual_data if isinstance(item, dict))
    return items


def load_help_kb_items() -> list[dict[str, Any]]:
    """Load help知识库.docx as article-level retrieval chunks."""
    global _HELP_KB_CACHE, _HELP_KB_MTIME_NS

    try:
        mtime_ns = HELP_KB_PATH.stat().st_mtime_ns
    except OSError:
        return []
    if _HELP_KB_CACHE and _HELP_KB_MTIME_NS == mtime_ns:
        return _HELP_KB_CACHE

    try:
        from docx import Document
    except ImportError:
        return []

    try:
        document = Document(HELP_KB_PATH)
    except (OSError, ValueError):
        return []

    chunks: list[dict[str, Any]] = []
    current_title = ""
    current_source = ""
    current_lines: list[str] = []

    def flush_chunk() -> None:
        nonlocal current_title, current_source, current_lines
        body = "\n".join(line for line in current_lines if line).strip()
        if not current_title or not body:
            current_lines = []
            return
        chunks.append(
            {
                "id": f"help-kb-{len(chunks) + 1}",
                "category": "help知识库",
                "question": current_title,
                "summary": body[:1200],
                "keywords": [current_title, current_source],
                "answer": {
                    "entry": current_source,
                    "steps": [body[:5000]],
                    "check": "",
                    "risk": "",
                },
                "source": current_source,
                "knowledgeText": body[:5000],
            }
        )
        current_lines = []

    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if not text:
            continue
        style_name = paragraph.style.name if paragraph.style else ""
        if style_name.startswith("Heading 1"):
            flush_chunk()
            current_title = text
            current_source = ""
            continue
        if style_name == "KB Source" or text.startswith("来源：http"):
            current_source = text.removeprefix("来源：").strip()
            continue
        if current_title:
            current_lines.append(text)

    flush_chunk()
    _HELP_KB_CACHE = chunks
    _HELP_KB_MTIME_NS = mtime_ns
    return chunks


def load_jsonl_kb_items(path: Path, source_name: str) -> list[dict[str, Any]]:
    try:
        mtime_ns = path.stat().st_mtime_ns
    except OSError:
        return []
    cache_key = str(path)
    cached = _JSONL_KB_CACHE.get(cache_key)
    if cached and cached[0] == mtime_ns:
        return cached[1]

    items: list[dict[str, Any]] = []
    try:
        lines = path.read_text(encoding="utf-8").splitlines()
    except OSError:
        return []

    for index, line in enumerate(lines, start=1):
        if not line.strip():
            continue
        try:
            raw = json.loads(line)
        except json.JSONDecodeError:
            continue
        if not isinstance(raw, dict):
            continue
        title = str(raw.get("question") or raw.get("title") or raw.get("path") or f"{source_name}-{index}")
        content = str(raw.get("knowledgeText") or raw.get("content") or raw.get("summary") or "")
        steps = raw.get("operation_steps") or (raw.get("answer") or {}).get("steps") or []
        notes = raw.get("permission_or_notes") or []
        if steps:
            content = (content + "\n步骤：" + "\n".join(str(step) for step in steps)).strip()
        if notes:
            content = (content + "\n注意：" + "\n".join(str(note) for note in notes)).strip()
        keywords = raw.get("retrieval_keywords") or raw.get("keywords") or []
        item = {
            "id": raw.get("id") or f"{source_name}-{index}",
            "category": raw.get("category") or source_name,
            "question": title,
            "summary": raw.get("summary") or content[:1200],
            "keywords": [title, raw.get("path") or "", *keywords],
            "answer": raw.get("answer")
            or {
                "entry": raw.get("path") or raw.get("source") or source_name,
                "steps": steps or [content[:5000]],
                "check": "",
                "risk": "涉及权限、发布范围、删除或考试设置时，提交前先核对影响对象。",
            },
            "source": raw.get("source") or source_name,
            "knowledgeText": content[:6000],
            "images": raw.get("images") or [],
        }
        items.append(item)

    _JSONL_KB_CACHE[cache_key] = (mtime_ns, items)
    return items


def load_precondition_items() -> list[dict[str, Any]]:
    global _PRECONDITION_CACHE, _PRECONDITION_MTIME_NS
    try:
        mtime_ns = PRECONDITION_PATH.stat().st_mtime_ns
    except OSError:
        return []
    if _PRECONDITION_CACHE and _PRECONDITION_MTIME_NS == mtime_ns:
        return _PRECONDITION_CACHE
    try:
        text = PRECONDITION_PATH.read_text(encoding="utf-8").strip()
    except OSError:
        return []
    if not text:
        _PRECONDITION_CACHE = []
        _PRECONDITION_MTIME_NS = mtime_ns
        return []
    sections = [
        section.strip()
        for section in re.split(r"(?m)^#{1,3}\s+|^\d+[.、]\s+", text)
        if section.strip()
    ]
    if not sections:
        sections = [text]
    items = []
    for index, section in enumerate(sections, start=1):
        title = section.splitlines()[0][:80]
        items.append(
            {
                "id": f"incomplete-precondition-{index}",
                "category": "不完全前置条件",
                "question": title,
                "summary": section[:1200],
                "keywords": [title, "前置条件", "不完全前置条件", "条件不足", "缺少条件"],
                "answer": {
                    "entry": "前置条件说明",
                    "steps": [section[:5000]],
                    "check": "",
                    "risk": "当前操作可能存在前置条件不完整的情况，回答时优先提醒用户补充页面、角色或操作对象。",
                },
                "source": PRECONDITION_PATH.name,
                "knowledgeText": section[:5000],
            }
        )
    _PRECONDITION_CACHE = items
    _PRECONDITION_MTIME_NS = mtime_ns
    return items


def load_extended_kb_items() -> list[dict[str, Any]]:
    """Legacy extended knowledge is disabled in FAQ-only text mode."""
    return []


QUERY_EXPANSIONS = {
    "ai": ["AI助手", "AI助教", "AI工作台", "教师备课助手", "智能助手"],
    "智能体": ["AI助手", "扣子智能体", "新增AI助手", "创建AI助手"],
    "助手": ["AI助手", "AI助教", "教师备课助手", "AI工作台"],
    "公告": ["新建公告", "编辑公告", "删除公告", "发布公告"],
    "发公告": ["发布公告", "新建公告", "课程公告"],
    "发布公告": ["新建公告", "课程公告"],
    "新建公告": ["发布新公告", "课程公告"],
    "课件": ["新建课件", "导入课件", "关联课件", "课件库"],
    "资源": ["添加资源", "教学资源", "资源检索", "课堂中打开资源"],
    "班级": ["创建班级", "批量导入班级", "学生入班", "班级管理"],
    "学生": ["加入学生", "导入学生", "学生扫码", "学生画像"],
    "考试": ["添加考试", "发布考试", "试卷", "题库"],
    "题库": ["新建题库文件夹", "试题库", "智能出题"],
    "ppt": ["生成PPT", "主题生成PPT", "大纲生成PPT"],
    "银行账号": ["账号", "个人资料", "修改个人资料", "个人信息"],
    "账号": ["个人资料", "个人信息", "密码", "登录"],
}


ACTION_GROUPS = {
    "公告": {
        "新建": ["新建", "新增", "创建", "发布", "发"],
        "编辑": ["编辑", "修改"],
        "删除": ["删除", "移除", "撤"],
    },
    "课程": {
        "创建": ["创建", "新建", "新增", "建课"],
        "删除": ["删除", "移除"],
    },
    "AI助手": {
        "创建": ["创建", "新建"],
        "添加": ["添加", "新增", "选择"],
        "使用": ["使用", "打开"],
    },
}


def detected_action(question: str, object_name: str) -> str:
    text = normalize(question)
    if normalize(object_name) not in text:
        return ""
    for action, words in ACTION_GROUPS.get(object_name, {}).items():
        if any(normalize(word) in text for word in words):
            return action
    return ""


def action_mismatch_penalty(question: str, item: dict[str, Any]) -> int:
    title = normalize(
        " ".join(
            [
                str(item.get("question", "")),
                str(item.get("summary", "")),
                str(item.get("sourcePath", "")),
                str(item.get("path", "")),
            ]
        )
    )
    penalty = 0
    for object_name, actions in ACTION_GROUPS.items():
        wanted = detected_action(question, object_name)
        if not wanted:
            continue
        for action, words in actions.items():
            if action == wanted:
                continue
            if any(normalize(word) in title for word in words):
                penalty += 45
    return penalty


def expanded_question_text(question: str) -> str:
    normalized = normalize(question)
    additions: list[str] = []
    for trigger, values in QUERY_EXPANSIONS.items():
        normalized_trigger = normalize(trigger)
        if normalized_trigger == "公告" and detected_action(question, "公告"):
            continue
        if normalized_trigger in normalized:
            additions.extend(values)
    if not additions:
        return question
    return question + " " + " ".join(additions)


def intent_core_text(question: str) -> str:
    return canonical_intent_text(question)


def help_kb_score(question: str, item: dict[str, Any]) -> int:
    expanded_question = expanded_question_text(question)
    query = normalize(expanded_question)
    raw_query = normalize(question)
    core_query = intent_core_text(question)
    title = normalize(item.get("question", ""))
    body = normalize(item.get("knowledgeText", ""))
    if not query:
        return 0

    score = 0
    if raw_query and raw_query == title:
        score += 70
    elif raw_query and (raw_query in title or title in raw_query):
        score += 45
    if core_query and len(core_query) >= 3:
        if core_query == title:
            score += 80
        elif core_query in title or title in core_query:
            score += 60
    if query == title:
        score += 40
    elif query in title or title in query:
        score += 20

    query_parts = [
        part
        for part in re.split(r"[\s，。！？、；：/()（）【】<>]+", expanded_question.lower())
        if len(part) >= 2
    ]
    for part in query_parts:
        normalized_part = normalize(part)
        if normalized_part in title:
            score += 8
        elif normalized_part in body:
            score += 3

    query_bigrams = {query[index : index + 2] for index in range(len(query) - 1)}
    title_bigrams = {title[index : index + 2] for index in range(len(title) - 1)}
    body_bigrams = {body[index : index + 2] for index in range(min(len(body) - 1, 1200))}
    score += len(query_bigrams & title_bigrams) * 3
    score += min(len(query_bigrams & body_bigrams), 8)
    return max(0, score - action_mismatch_penalty(question, item))


def find_help_kb_matches(
    question: str,
    items: list[dict[str, Any]],
    limit: int = 2,
) -> list[dict[str, Any]]:
    matches = [{"item": item, "score": help_kb_score(question, item)} for item in items]
    matches = [match for match in matches if match["score"] >= 6]
    matches.sort(key=lambda match: match["score"], reverse=True)
    if not matches:
        return []
    best_score = matches[0]["score"]
    return [match for match in matches if match["score"] >= best_score * 0.65][:limit]


def normalize(value: Any) -> str:
    return re.sub(r"[\s，。！？、；：/()（）【】<>“”‘’\"']+", "", str(value).lower())


def canonical_intent_text(value: Any) -> str:
    """Normalize harmless wording differences without merging different functions."""
    text = normalize(value)
    for prefix in ["麻烦", "请问", "请教", "请", "我想要", "我想", "我要", "想要", "能不能", "可不可以"]:
        if text.startswith(prefix):
            text = text[len(prefix) :]
            break
    for prefix in ["怎么", "怎样", "如何"]:
        if text.startswith(prefix):
            text = text[len(prefix) :]
            break
    text = re.sub(r"^(快速|立即|直接|马上|迅速|便捷)+", "", text)
    text = re.sub(r"^(新建|新增)", "创建", text)
    return re.sub(r"^(创建|新建|新增|添加|做|建)(?:一个|一门|一项|一份|一名|个)", r"\1", text)


def split_term(term: str) -> list[str]:
    return [part for part in re.split(r"[，。！？、；：\s/()（）【】<>]+", term) if part]


def searchable_terms(item: dict[str, Any]) -> list[str]:
    terms = [
        item.get("question", ""),
        item.get("summary", ""),
        item.get("category", ""),
    ]
    terms.extend(item.get("keywords") or [])
    answer = item.get("answer") or {}
    terms.append(answer.get("entry", ""))
    terms.extend(answer.get("steps") or [])
    return [normalize(term) for term in terms if term]


def score_item(question: str, item: dict[str, Any]) -> int:
    raw_text = normalize(question)
    core_text = canonical_intent_text(question)
    best_score = 0
    supporting_hits = 0
    for term in searchable_terms(item):
        if not term:
            continue
        canonical_term = canonical_intent_text(term)
        if raw_text and raw_text == term:
            term_score = 240
        else:
            if core_text and len(core_text) >= 3 and core_text == canonical_term:
                # Exact intent equivalents must tie with the literal wording so
                # harmless modifiers do not switch to another source record.
                term_score = 240
            elif core_text and len(core_text) >= 4 and (
                core_text in canonical_term or canonical_term in core_text
            ):
                coverage = min(len(core_text), len(canonical_term)) / max(
                    len(core_text), len(canonical_term)
                )
                term_score = 30 + round(coverage * 40)
            else:
                term_score = 0
        for part in split_term(term):
            normalized_part = canonical_intent_text(part)
            if len(normalized_part) >= 3 and normalized_part in core_text:
                supporting_hits += 1
        best_score = max(best_score, term_score)
    score = best_score + min(supporting_hits, 5)
    return max(0, score - action_mismatch_penalty(question, item))


def find_matches(question: str, items: list[dict[str, Any]], limit: int = 3) -> list[dict[str, Any]]:
    matches = [{"item": item, "score": score_item(question, item)} for item in items]
    matches = [match for match in matches if match["score"] >= 35]
    matches.sort(key=lambda match: match["score"], reverse=True)
    return matches[:limit]


WORKFLOW_RELATED_IDS = {
    "course-create": ["teaching-team-full-flow", "class-setup", "courseware-create"],
    "teaching-team": ["class-setup", "invite-students", "student-batch-import"],
    "class-setup": ["invite-students", "teaching-team", "student-batch-import"],
    "invite-students": ["student-batch-import", "teaching-team", "courseware-link"],
    "courseware-link": ["courseware-create", "courseware-import-library", "class-start"],
    "resource-add": ["courseware-import-local", "courseware-link", "class-start"],
    "personal-homework": ["peer-review", "group-homework", "progress-score"],
    "quiz-publish": ["question-bank-add", "class-start", "progress-score"],
    "exam-arrange": ["paper-manual", "exam-grade", "exam-result"],
    "exam-result": ["exam-analysis", "exam-grade", "assessment-rule"],
}


def item_term_set(item: dict[str, Any]) -> set[str]:
    terms: set[str] = set()
    for term in searchable_terms(item):
        for part in split_term(term):
            normalized = normalize(part)
            if len(normalized) >= 2:
                terms.add(normalized)
    return terms


def related_item_score(source: dict[str, Any], candidate: dict[str, Any], question: str) -> int:
    score = 0
    if source.get("category") and source.get("category") == candidate.get("category"):
        score += 5
    overlap = item_term_set(source) & item_term_set(candidate)
    score += min(len(overlap), 5) * 2
    question_score = score_item(question, candidate)
    if question_score:
        score += min(question_score, 8)
    return score


def recommend_related_questions(
    question: str,
    matched_items: list[dict[str, Any]],
    faq_items: list[dict[str, Any]],
    limit: int = 3,
) -> list[dict[str, str]]:
    by_id = {str(item.get("id")): item for item in faq_items if item.get("id")}
    excluded_ids = {str(item.get("id")) for item in matched_items if item.get("id")}
    excluded_questions = {
        canonical_intent_text(item.get("question", ""))
        for item in matched_items
    }
    ranked: dict[str, tuple[int, dict[str, Any]]] = {}

    def add(item_id: str, score: int) -> None:
        item = by_id.get(item_id)
        if not item or item_id in excluded_ids:
            return
        item_question = str(item.get("question", "")).strip()
        if (
            not item_question
            or canonical_intent_text(item_question) in excluded_questions
        ):
            return
        previous = ranked.get(item_id)
        if previous is None or score > previous[0]:
            ranked[item_id] = (score, item)

    for index, source in enumerate(matched_items):
        base = 100 - index * 10
        for position, item_id in enumerate(source.get("related") or []):
            add(str(item_id), base - position)
        for position, item_id in enumerate(WORKFLOW_RELATED_IDS.get(str(source.get("id")), [])):
            add(item_id, base - 10 - position)

    for candidate in faq_items:
        candidate_id = str(candidate.get("id") or "")
        if not candidate_id or candidate_id in excluded_ids:
            continue
        similarity = max(
            (related_item_score(source, candidate, question) for source in matched_items),
            default=score_item(question, candidate),
        )
        if similarity >= 5:
            add(candidate_id, similarity)

    sorted_items = sorted(
        ranked.values(),
        key=lambda entry: (-entry[0], str(entry[1].get("question", ""))),
    )
    recommendations: list[dict[str, str]] = []
    recommended_intents: set[str] = set()
    for _, item in sorted_items:
        item_question = str(item.get("question", ""))
        intent = canonical_intent_text(item_question)
        if not intent or intent in excluded_questions or intent in recommended_intents:
            continue
        recommended_intents.add(intent)
        recommendations.append(
            {
                "label": item_question.rstrip("？?"),
                "question": item_question,
            }
        )
        if len(recommendations) >= limit:
            break
    return recommendations


def is_greeting_or_small_talk(question: str) -> bool:
    text = normalize(question)
    exact_greetings = {
        "你好",
        "您好",
        "hi",
        "hello",
        "在吗",
        "早上好",
        "下午好",
        "晚上好",
        "谢谢",
        "感谢",
        "你是谁",
        "你能做什么",
        "ok",
        "okk",
        "okkk",
        "好的",
        "好",
        "嗯",
        "嗯嗯",
        "收到",
        "明白",
        "可以",
        "行",
    }
    fuzzy_greetings = {"你好", "您好", "hi", "hello", "在吗", "谢谢", "感谢", "你是谁", "你能做什么"}
    return text in exact_greetings or (len(text) <= 12 and any(word in text for word in fuzzy_greetings))


def small_talk_answer(question: str) -> str:
    text = normalize(question)
    if "谢谢" in text or "感谢" in text:
        return "不客气。你可以继续问我 uLearning Web 端教师操作问题，比如作业、考试、班级或课程设置。"
    if "你是谁" in text or "你能做什么" in text:
        return "我是 AI操作助手，主要帮老师定位 uLearning Web 端的功能入口，并把操作步骤说明清楚。"
    if text in {"ok", "okk", "okkk", "好的", "好", "嗯", "嗯嗯", "收到", "明白", "可以", "行"}:
        return "好的。你可以继续问我 uLearning Web 端教师操作问题。"
    return "你好，我是 AI操作助手。你可以直接问我 uLearning Web 端教师操作问题，比如怎么布置作业、安排考试或设置班级。"


def is_scope_question(question: str) -> bool:
    text = normalize(question)
    patterns = [
        "只能问",
        "可以问什么",
        "能问什么",
        "问你什么",
        "哪些问题",
        "什么问题",
        "问题库",
        "知识库",
        "范围",
    ]
    return any(pattern in text for pattern in patterns) and (
        "问" in text or "问题" in text or "范围" in text or "知识库" in text
    )


def scope_answer() -> str:
    return "\n".join(
        [
            "不是只能问固定问题。你可以直接用自己的话问，我会尽量帮你定位到对应操作。",
            "",
            "比较适合问这些方面：",
            "课程：创建课程、课程封面、教学团队、班级设置",
            "教学：课件、资源、学习计划、学习进度",
            "活动：作业、互评、小组作业、测验、讨论",
            "考试：题库、组卷、安排考试、阅卷、成绩分析",
            "成绩：考核规则、学习时长、作业成绩、课堂点名",
            "异常：按钮找不到、权限不足、保存失败、系统报错",
            "",
            "如果暂时没有收录，我也会记录下来，后续可以补充进 FAQ。",
        ]
    )


TROUBLESHOOTING_TRIGGERS = (
    "看不到",
    "找不到",
    "没看到",
    "没有看到",
    "没有显示",
    "不显示",
    "无法",
    "不能",
    "失败",
    "打不开",
    "不见了",
    "没反应",
    "报错",
)

TROUBLESHOOTING_RULES = [
    {
        "id": "homework-not-visible",
        "objects": ("作业",),
        "questions": [
            "这份作业是否已经点击【发布】，而不是仍保存在草稿中？",
            "发布时选择的班级里，是否包含这名学生所在的班级？",
            "这名学生是否已经加入当前课程的正确班级？",
            "作业的开始时间是否已经到达，截止时间是否尚未结束？",
            "学生登录的是否是正确账号，并且查看的是当前学期课程？",
        ],
        "next": "先确认以上五项。哪一项不确定，就从那一项开始检查；如果全部满足，再提供作业名称、班级和学生端当前页面。",
    },
    {
        "id": "exam-not-visible",
        "objects": ("考试", "测验"),
        "questions": [
            "考试或测验是否已经发布？",
            "发布对象是否包含该学生所在班级？",
            "学生是否已加入当前课程的正确班级？",
            "考试的开放时间是否已经到达，结束时间是否尚未超过？",
            "学生登录的账号和所在机构是否正确？",
        ],
        "next": "先核对发布状态、班级和时间。全部正确仍看不到时，请提供考试名称、班级和学生端页面。",
    },
    {
        "id": "courseware-not-visible",
        "objects": ("课件", "章节", "单元"),
        "questions": [
            "课件是否已经保存并发布？",
            "课件是否已经关联到学生所在的班课？",
            "学生是否已加入当前课程的正确班级？",
            "学习计划或章节开放时间是否允许当前查看？",
            "学生查看的是否是当前课程和当前学期？",
        ],
        "next": "优先检查“是否发布”和“是否关联班课”。这两项满足后，再检查开放时间与学生账号。",
    },
    {
        "id": "resource-not-visible",
        "objects": ("资源", "视频", "文档", "文件"),
        "questions": [
            "资源是否已经上传完成并保存？",
            "视频或文件是否仍处于转码、处理中？",
            "资源是否已经添加到学生可见的课程或单元中？",
            "可见班级是否包含该学生所在班级？",
            "当前问题出现在电脑端还是 App？",
        ],
        "next": "如果资源仍在处理，请等待处理完成；否则继续核对资源所在单元和可见班级。",
    },
    {
        "id": "announcement-not-visible",
        "objects": ("公告", "通知"),
        "questions": [
            "公告是否已经正式发布？",
            "发布对象是否包含该学生所在班级？",
            "学生是否已加入当前课程的正确班级？",
            "学生查看的是否是当前课程的公告页面？",
        ],
        "next": "先检查公告发布状态和发布班级；两项都正确时，再核对学生账号和课程。",
    },
    {
        "id": "entry-or-button-missing",
        "objects": ("按钮", "入口", "功能", "设置"),
        "questions": [
            "当前账号角色是课程管理员、任课教师还是助教？",
            "当前是否进入了正确课程和正确功能页面？",
            "该账号是否已被分配到对应班级？",
            "当前使用的是电脑端还是 App？",
        ],
        "next": "按钮缺失通常与账号角色、所在页面或端类型有关。请先提供当前页面名称和账号角色。",
    },
    {
        "id": "student-course-missing",
        "objects": ("班课", "课程"),
        "questions": [
            "学生登录的是否是正确账号和学校机构？",
            "学生是否已经通过班级编码或二维码加入班级？",
            "教师查看的班级名单中是否能找到该学生？",
            "学生查看的是否是当前学期课程？",
        ],
        "next": "如果教师名单中没有该学生，请先让学生加入正确班级；名单中已有学生时，再检查账号和学期。",
    },
]


def get_troubleshooting_rule(question: str) -> dict[str, Any] | None:
    text = normalize(question)
    has_issue_trigger = any(
        normalize(trigger) in text for trigger in TROUBLESHOOTING_TRIGGERS
    )
    has_missing_ui_pattern = bool(
        re.search(r"(?:没有|没找到|未显示).*(?:按钮|入口|功能|设置)", text)
    )
    if not has_issue_trigger and not has_missing_ui_pattern:
        return None
    for rule in TROUBLESHOOTING_RULES:
        if any(normalize(object_name) in text for object_name in rule["objects"]):
            return rule
    return {
        "id": "general-operation-failure",
        "questions": [
            "当前操作是否已经保存或发布成功？",
            "操作对象和可见班级是否选择正确？",
            "当前账号角色是否有该功能权限？",
            "页面显示的时间范围或状态是否允许当前操作？",
            "刷新页面或重新登录后，问题是否仍然存在？",
        ],
        "next": "请补充当前页面名称、账号角色、操作对象和页面提示，我再帮你定位到具体环节。",
    }


def troubleshooting_answer(rule: dict[str, Any]) -> str:
    lines = [
        "这类情况先不要重复创建或重新发布，我们先定位是哪一个前置条件没有满足。",
        "",
        "请依次确认：",
    ]
    lines.extend(
        f"{index}、{question}"
        for index, question in enumerate(rule.get("questions") or [], start=1)
    )
    lines.extend(["", f"下一步：{rule.get('next')}"])
    return "\n".join(lines)


PROMPT_GUIDE_IDS = {
    "overview": [
        "course-create",
        "teaching-team",
        "class-setup",
        "courseware-link",
        "learning-plan",
        "resource-add",
        "personal-homework",
        "quiz-publish",
        "discussion",
        "exam-arrange",
        "assessment-rule",
        "progress-score",
    ],
    "start": [
        "teacher-quick-start-text",
    ],
    "course": [
        "course-create",
        "teaching-team",
        "class-setup",
        "invite-students",
        "courseware-link",
        "learning-plan",
    ],
}


def get_prompt_guide_items(question: str, items: list[dict[str, Any]]) -> list[dict[str, Any]]:
    text = normalize(question)
    overview_triggers = [
        "这个平台可以做什么",
        "ulearning有哪些功能",
        "平台有哪些功能",
        "有哪些功能",
        "能做什么",
    ]
    start_triggers = [
        "我应该从哪里开始",
        "第一次使用",
        "应该先做什么",
        "从哪里开始",
    ]
    guide_key = None
    if any(trigger in text for trigger in overview_triggers):
        guide_key = "overview"
    elif any(trigger in text for trigger in start_triggers):
        guide_key = "start"

    if not guide_key:
        return []

    by_id = {item.get("id"): item for item in items}
    return [by_id[item_id] for item_id in PROMPT_GUIDE_IDS[guide_key] if item_id in by_id]


def is_probably_ulearning_related(question: str) -> bool:
    text = normalize(question)
    platform_terms = [
        "ulearning",
        "ul",
        "课程",
        "班级",
        "班课",
        "学生",
        "老师",
        "教师",
        "助教",
        "教学",
        "课件",
        "资源",
        "作业",
        "互评",
        "测验",
        "讨论",
        "考试",
        "试卷",
        "题库",
        "试题",
        "组卷",
        "成绩",
        "考核",
        "学习计划",
        "学习进度",
        "学习时长",
        "点名",
        "签到",
        "公告",
        "权限",
        "按钮",
        "入口",
        "页面",
        "发布",
        "保存",
        "导入",
        "上传",
        "报错",
        "看不到",
        "找不到",
        "登录",
        "账号",
        "个人信息",
        "ai助手",
        "ai助教",
        "学生ai助教",
        "教师备课助手",
        "备课助手",
        "智能助手",
        "课程工作台",
    ]
    return any(term in text for term in platform_terms)


def unrelated_answer(question: str) -> str:
    return "\n".join(
        [
            "我是 AI操作助手，主要解答 uLearning Web 端教师操作问题。",
            "",
            "你可以问我这些方面：",
            "课程、班级、教学团队、课件资源",
            "作业、测验、讨论、考试、题库",
            "成绩、考核规则、权限、按钮找不到或系统报错",
            "",
            f"我无法解答“{question}”这类与 uLearning 教师操作无关的问题。",
        ]
    )


def record_unknown_question(question: str) -> None:
    record = {
        "question": question,
        "createdAt": datetime.now(timezone.utc).isoformat(),
        "status": "new",
    }
    with UNKNOWN_LOG_PATH.open("a", encoding="utf-8") as file:
        file.write(json.dumps(record, ensure_ascii=False) + "\n")


def fallback_answer(question: str) -> str:
    record_unknown_question(question)
    return "\n".join(
        [
            "这个问题我还需要结合页面继续定位，我会先记录下来，后续可以补充成更精确的操作指引。",
            "",
            "你可以先补充任意一项，我会按最接近的操作继续判断：",
            "当前页面名称",
            "你想操作的对象，比如课程、班级、作业、考试",
            "你看到的按钮、报错或权限提示",
            "",
            f"你的问题：{question}",
        ]
    )


def format_faq_answer(item: dict[str, Any]) -> str:
    answer = item.get("answer") or {}
    lines = []
    entry = answer.get("entry")
    if entry:
        lines.append(f"入口位置：{entry}")
    steps = answer.get("steps") or []
    if steps:
        cleaned_steps = []
        for step in steps:
            text = str(step).strip()
            text = re.split(r"\nTABLE\s+\d+|\n评分项\s*\|", text, maxsplit=1)[0]
            text = "\n".join(line for line in text.splitlines()[:6] if line.strip())
            if len(text) > 700:
                text = text[:700].rstrip() + "……"
            if text:
                cleaned_steps.append(text.strip("。"))
        if len(cleaned_steps) == 1:
            cleaned_steps.append("按上述说明处理后，返回当前页面检查结果是否已经生效")
        if cleaned_steps:
            lines.append("")
            lines.append("操作步骤：")
            lines.extend(f"{index}、{step}。" for index, step in enumerate(cleaned_steps, start=1))
    check = answer.get("check")
    risk = answer.get("risk")
    if not check:
        check = "返回当前模块的列表或详情页，确认刚才的操作结果已经显示。"
    if not risk:
        risk = "如果页面入口、按钮名称或结果与上述说明不一致，请先核对当前账号角色和所在页面。"
    lines.append("")
    lines.append(f"完成后看：{check}")
    lines.append("")
    lines.append(f"注意：{risk}")
    return "\n".join(lines)


def load_rewrite_prompt() -> str:
    try:
        return PROMPT_PATH.read_text(encoding="utf-8")
    except OSError:
        return (
            "你是 uLearning Web 端教师专用的操作指南助手。"
            "可以自然处理寒暄；操作问题只基于参考答案回答。"
            "先说明入口位置，再说明操作路径、完成后检查和注意事项。"
        )


def build_rewrite_messages(
    question: str,
    items: list[dict[str, Any]],
    related_questions: list[dict[str, str]] | None = None,
) -> list[dict[str, str]]:
    return [
        {
            "role": "system",
            "content": load_rewrite_prompt(),
        },
        {
            "role": "user",
            "content": json.dumps(
                {
                    "teacher_question": question,
                    "reference_answers": items,
                    "candidate_followup_questions": related_questions or [],
                    "required_style": (
                        "先主动判断用户真实想完成的操作，再比较各FAQ与问题的关联度。"
                        "回答必须以关联度最高且能直接帮助当前操作的FAQ为依据。"
                        "候选后续问题由页面按钮展示，不要把它们逐条重复写进正文。"
                        "如果用户一次问多个操作，请把多个命中的参考答案自然串接回答。"
                        "先用一句话说明用户要完成什么，再定位入口并写操作路径。"
                        "只回答用户当前要做的这一件事，不要主动展开后续所有分支。"
                        "所有操作问题都必须使用“入口位置、操作步骤、完成后看、注意”的固定结构。"
                        "操作步骤必须用1、2、3这样的阿拉伯数字编号，写出完整且可照做的步骤，不得省略为概括描述。"
                        "每一步只写一个动作，按钮名使用【】突出；不要把多个点击动作塞进一个长句。"
                        "涉及发布、班级、时间、成绩或删除时，最后单独写注意事项。"
                        "回答要让第一次使用平台的老师也能照着完成，不能只给概括性结论。"
                        "遇到多个操作方向时只列选项名称，让用户选择后再展开。"
                        "输出清楚优先于过度简短，适合 uLearning Web 端教师逐步照着操作。"
                        "不要使用 Markdown 加粗、星号、井号标题或表格。"
                        "如果用户问平台能做什么、有哪些功能、从哪里开始，只列举可做的功能名称，"
                        "按模块分组即可，不要在每个功能后追加解释性长句。"
                    ),
                },
                ensure_ascii=False,
            ),
        },
    ]


def call_deepseek_rewrite(
    question: str,
    items: list[dict[str, Any]],
    related_questions: list[dict[str, str]] | None = None,
) -> str:
    api_key = (
        os.getenv("DEEPSEEK_API_KEY")
        or os.getenv("ULEARNING_TEACHER_ASSISTANT_API_KEY")
        or os.getenv("OPENAI_API_KEY")
    )
    if not api_key:
        raise RuntimeError("Missing DEEPSEEK_API_KEY")

    base_url = os.getenv("DEEPSEEK_BASE_URL", "https://api.deepseek.com/v1").rstrip("/")
    model = os.getenv("DEEPSEEK_MODEL", "deepseek-chat")
    payload = {
        "model": model,
        "messages": build_rewrite_messages(question, items, related_questions),
        "temperature": 0.2,
        "stream": False,
    }
    request = urllib.request.Request(
        f"{base_url}/chat/completions",
        data=json.dumps(payload).encode("utf-8"),
        headers={
            "Authorization": f"Bearer {api_key}",
            "Content-Type": "application/json",
        },
        method="POST",
    )

    try:
        with urllib.request.urlopen(request, timeout=60) as response:
            data = json.loads(response.read().decode("utf-8"))
    except urllib.error.HTTPError as exc:
        body = exc.read().decode("utf-8", errors="replace")
        raise RuntimeError(f"DeepSeek request failed with HTTP {exc.code}: {body}") from exc

    choices = data.get("choices") or []
    if choices:
        content = ((choices[0].get("message") or {}).get("content") or "").strip()
        if content:
            return clean_model_answer(content)
    return "\n\n".join(format_faq_answer(item) for item in items)


GENERAL_CHAT_SUFFIX = "关于 uLearning，你还想做点什么吗？"


def call_deepseek_general_chat(question: str) -> str:
    api_key = (
        os.getenv("DEEPSEEK_API_KEY")
        or os.getenv("ULEARNING_TEACHER_ASSISTANT_API_KEY")
        or os.getenv("OPENAI_API_KEY")
    )
    if not api_key:
        raise RuntimeError("Missing DEEPSEEK_API_KEY")

    base_url = os.getenv("DEEPSEEK_BASE_URL", "https://api.deepseek.com/v1").rstrip("/")
    model = os.getenv("DEEPSEEK_MODEL", "deepseek-chat")
    payload = {
        "model": model,
        "messages": [
            {
                "role": "system",
                "content": (
                    "你是“小蜜蜂”，一个以 uLearning 教师操作指导为主要职责的智能助手。"
                    "用户闲聊、表达情绪或询问一般知识时，可以像通用 DeepSeek 助手一样自然、"
                    "灵活且有帮助地回答，不要生硬拒绝，也不要把普通闲聊误判成平台故障。"
                    "回答当前问题本身，保持简洁、友好、真实；不确定的事实不要编造。"
                    "不要声称已经替用户完成现实操作。"
                    f"回答结束后另起一行，固定补充：{GENERAL_CHAT_SUFFIX}"
                ),
            },
            {"role": "user", "content": question},
        ],
        "temperature": 0.7,
        "stream": False,
    }
    request = urllib.request.Request(
        f"{base_url}/chat/completions",
        data=json.dumps(payload).encode("utf-8"),
        headers={
            "Authorization": f"Bearer {api_key}",
            "Content-Type": "application/json",
        },
        method="POST",
    )
    try:
        with urllib.request.urlopen(request, timeout=60) as response:
            data = json.loads(response.read().decode("utf-8"))
    except urllib.error.HTTPError as exc:
        body = exc.read().decode("utf-8", errors="replace")
        raise RuntimeError(f"DeepSeek request failed with HTTP {exc.code}: {body}") from exc

    choices = data.get("choices") or []
    content = (
        ((choices[0].get("message") or {}).get("content") or "").strip()
        if choices
        else ""
    )
    if not content:
        raise RuntimeError("DeepSeek returned an empty response")
    cleaned = clean_model_answer(content)
    if GENERAL_CHAT_SUFFIX not in cleaned:
        cleaned = f"{cleaned}\n\n{GENERAL_CHAT_SUFFIX}"
    return cleaned


def general_chat_fallback(question: str) -> str:
    text = normalize(question)
    if "为什么不回答我" in text or "怎么不回答" in text:
        lead = "抱歉，刚才可能把你的话误判成了平台故障。现在你可以正常和我聊天，我也会继续帮你处理 uLearning 操作问题。"
    elif is_greeting_or_small_talk(question):
        lead = small_talk_answer(question)
    else:
        lead = "这个问题我暂时没能连接到通用问答服务。你可以稍后再问一次，我不会把普通聊天当成平台故障。"
    return f"{lead}\n\n{GENERAL_CHAT_SUFFIX}"


ENGLISH_QUESTION_ALIASES = {
    "quick start": "快速开始",
    "how do i create a course?": "如何创建课程？",
    "how do i set up a teaching team?": "如何设置教学团队？",
    "how do i create courseware?": "如何新建课件？",
    "how do i assign individual homework?": "如何布置个人作业？",
    "how do i start a class?": "如何开始上课？",
    "how do i view learning progress and grades?": "如何查看进度成绩？",
    "what can i do in ulearning?": "uLearning 有哪些功能？",
    "how do i set up preview chapters for courseware?": "如何设置课件试听章节？",
    "how do i set up trial chapters for courseware?": "如何设置课件试听章节？",
}


def deepseek_translate(text: str, target_language: str, purpose: str = "") -> str:
    api_key = (
        os.getenv("DEEPSEEK_API_KEY")
        or os.getenv("ULEARNING_TEACHER_ASSISTANT_API_KEY")
        or os.getenv("OPENAI_API_KEY")
    )
    if not api_key:
        raise RuntimeError("Missing DEEPSEEK_API_KEY")

    base_url = os.getenv("DEEPSEEK_BASE_URL", "https://api.deepseek.com/v1").rstrip("/")
    model = os.getenv("DEEPSEEK_MODEL", "deepseek-chat")
    system_prompt = (
        f"Translate the supplied text into {target_language}. "
        "Preserve the original meaning, numbered steps, paragraph breaks, product names, "
        "button labels, and the name uLearning. Use concise, natural wording suitable for "
        "a teacher-facing software help center. Do not add explanations or Markdown fences."
    )
    if purpose:
        system_prompt += f" Context: {purpose}"
    payload = {
        "model": model,
        "messages": [
            {"role": "system", "content": system_prompt},
            {"role": "user", "content": text},
        ],
        "temperature": 0.1,
        "stream": False,
    }
    request = urllib.request.Request(
        f"{base_url}/chat/completions",
        data=json.dumps(payload).encode("utf-8"),
        headers={
            "Authorization": f"Bearer {api_key}",
            "Content-Type": "application/json",
        },
        method="POST",
    )
    with urllib.request.urlopen(request, timeout=60) as response:
        data = json.loads(response.read().decode("utf-8"))
    choices = data.get("choices") or []
    translated = (
        ((choices[0].get("message") or {}).get("content") or "").strip()
        if choices
        else ""
    )
    if not translated:
        raise RuntimeError("DeepSeek returned an empty translation")
    return clean_model_answer(translated)


def resolve_english_question(question: str) -> str:
    alias = ENGLISH_QUESTION_ALIASES.get(question.strip().lower())
    if alias:
        return alias
    normalized_question = normalize(question)
    try:
        manual_guides = json.loads(MANUAL_2026_GUIDES_PATH.read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError):
        manual_guides = []
    if isinstance(manual_guides, list):
        for guide in manual_guides:
            if not isinstance(guide, dict):
                continue
            english_candidates = [guide.get("questionEn"), *(guide.get("aliases") or [])]
            if any(normalize(str(candidate)) == normalized_question for candidate in english_candidates if candidate):
                return str(guide.get("question") or question)
    return deepseek_translate(
        question,
        "Simplified Chinese",
        "Convert the user's English uLearning help question into a concise Chinese search query.",
    )


def english_translation_fallback(result: dict[str, Any]) -> str:
    if result.get("intent") == "troubleshooting":
        return (
            "Let’s identify which prerequisite is missing before repeating the operation.\n\n"
            "1. Confirm that the item was saved or published successfully.\n"
            "2. Check that the correct course, class, and target users were selected.\n"
            "3. Confirm that your account has permission for this operation.\n"
            "4. Check the active time range and current status.\n"
            "5. Refresh the page or sign in again, then test once more.\n\n"
            "If the issue remains, share the page name, your role, the target item, and the on-screen message."
        )
    return (
        "I found the relevant uLearning guide, but the English translation service is "
        "temporarily unavailable. Please try again in a moment."
    )


def answer_question_localized(question: str, language: str = "zh") -> dict[str, Any]:
    if language.lower() not in {"en", "english", "en-us", "en-gb"}:
        return answer_question(question)

    try:
        resolved_question = (
            resolve_english_question(question)
            if re.search(r"[A-Za-z]", question) and not re.search(r"[\u4e00-\u9fff]", question)
            else question
        )
    except Exception:
        resolved_question = ENGLISH_QUESTION_ALIASES.get(
            question.strip().lower(),
            question,
        )

    result = answer_question(resolved_question)
    try:
        result["answer"] = deepseek_translate(
            str(result.get("answer") or ""),
            "English",
            "Translate a structured uLearning help answer. Keep every numbered step.",
        )
        related = result.get("relatedQuestions") or []
        if related:
            translated_related = []
            for item in related[:4]:
                translated_question = deepseek_translate(
                    str(item.get("question") or ""),
                    "English",
                    "Translate a short suggested follow-up question.",
                )
                translated_related.append(
                    {
                        "label": translated_question.rstrip("?."),
                        "question": translated_question,
                    }
                )
            result["relatedQuestions"] = translated_related
    except Exception:
        result["answer"] = english_translation_fallback(result)
        result["relatedQuestions"] = []

    result["resolvedQuestion"] = resolved_question
    result["language"] = "en"
    return result


def clean_model_answer(answer: str) -> str:
    return (
        answer.replace("根据FAQ，", "")
        .replace("根据 FAQ，", "")
        .replace("根据FAQ", "")
        .replace("根据 FAQ", "")
        .replace("根据参考答案，", "")
        .replace("根据参考答案", "")
        .replace("根据资料，", "")
        .replace("根据资料", "")
        .replace("**", "")
        .replace("__", "")
        .replace("###", "")
        .replace("##", "")
        .replace("#", "")
        .strip()
    )


def clean_prompt_guide_answer(answer: str) -> str:
    cleaned_lines = []
    for raw_line in clean_model_answer(answer).splitlines():
        line = raw_line.strip()
        if not line:
            cleaned_lines.append("")
            continue
        if "告诉我" in line or "你想" in line or "可以直接" in line:
            continue
        if line.endswith("？") or line.endswith("?"):
            continue
        cleaned_lines.append(raw_line)
    return "\n".join(cleaned_lines).strip()


def answer_question(question: str) -> dict[str, Any]:
    faq_items = load_faq_items()
    help_kb_items: list[dict[str, Any]] = []
    prompt_guide_items = get_prompt_guide_items(question, faq_items)
    if prompt_guide_items:
        related_questions = recommend_related_questions(question, prompt_guide_items, faq_items)
        try:
            answer = clean_prompt_guide_answer(
                call_deepseek_rewrite(question, prompt_guide_items, related_questions)
            )
            model_used = True
        except Exception:
            answer = "\n\n".join(format_faq_answer(item) for item in prompt_guide_items)
            model_used = False
        return {
            "answer": answer,
            "matched": True,
            "modelUsed": model_used,
            "intent": "prompt_guide",
            "relatedQuestions": related_questions,
            "matches": [
                {
                    "id": item.get("id"),
                    "question": item.get("question"),
                    "score": 10,
                }
                for item in prompt_guide_items
            ],
        }

    if is_greeting_or_small_talk(question):
        try:
            answer = call_deepseek_general_chat(question)
            model_used = True
        except Exception:
            answer = general_chat_fallback(question)
            model_used = False
        return {
            "answer": answer,
            "matched": False,
            "modelUsed": model_used,
            "intent": "small_talk",
        }

    if is_scope_question(question):
        return {
            "answer": scope_answer(),
            "matched": False,
            "modelUsed": False,
            "intent": "scope",
        }

    troubleshooting_rule = get_troubleshooting_rule(question)
    if troubleshooting_rule:
        return {
            "answer": troubleshooting_answer(troubleshooting_rule),
            "matched": True,
            "modelUsed": False,
            "intent": "troubleshooting",
            "diagnosticId": troubleshooting_rule.get("id"),
            "relatedQuestions": [],
            "matches": [],
        }

    faq_matches = find_matches(question, faq_items, limit=2)
    help_kb_matches = find_help_kb_matches(question, help_kb_items)
    matches = faq_matches + help_kb_matches
    matches.sort(key=lambda match: match["score"], reverse=True)
    if matches:
        relevance_floor = max(4, matches[0]["score"] * 0.5)
        matches = [match for match in matches if match["score"] >= relevance_floor]
    matches = matches[:3]
    if not matches:
        if not is_probably_ulearning_related(question):
            try:
                answer = call_deepseek_general_chat(question)
                model_used = True
            except Exception:
                answer = general_chat_fallback(question)
                model_used = False
            return {
                "answer": answer,
                "matched": False,
                "modelUsed": model_used,
                "intent": "general_chat",
            }
        return {
            "answer": fallback_answer(question),
            "matched": False,
            "modelUsed": False,
            "recorded": True,
        }

    # Use only the strongest retrieval result. Combining weaker candidates caused
    # the model to merge unrelated procedures (for example AI assistant creation
    # and the teacher lesson-preparation assistant).
    matches = matches[:1]
    matched_items = [matches[0]["item"]]
    related_questions = recommend_related_questions(question, matched_items, faq_items)
    answer = format_faq_answer(matched_items[0])
    model_used = False

    return {
        "answer": answer,
        "matched": True,
        "modelUsed": model_used,
        "relatedQuestions": related_questions,
        "matches": [
                {
                    "id": match["item"].get("id"),
                    "question": match["item"].get("question"),
                    "score": match["score"],
                    "source": match["item"].get("source") or match["item"].get("category"),
                    "images": match["item"].get("images") or [],
                }
                for match in matches
            ],
        }


def append_feedback(payload: dict[str, Any]) -> None:
    path = feedback_table_path(payload)
    workbook = load_feedback_workbook(path)
    sheet = workbook.active
    row = {
        "用户": str(payload.get("user", "教师")),
        "时间": str(payload.get("time") or datetime.now().strftime("%Y-%m-%d %H:%M:%S")),
        "问题": str(payload.get("question", "")),
        "记录ID": str(payload.get("feedbackId", "")),
    }
    sheet.append([row[field] for field in FEEDBACK_FIELDS])
    workbook.save(path)


def delete_feedback(payload: dict[str, Any]) -> bool:
    path = feedback_table_path(payload)
    feedback_id = str(payload.get("feedbackId", "")).strip()
    if not feedback_id or not path.exists():
        return False
    workbook = load_feedback_workbook(path)
    sheet = workbook.active
    id_column = len(FEEDBACK_FIELDS)
    for row_index in range(sheet.max_row, 1, -1):
        if str(sheet.cell(row=row_index, column=id_column).value or "") == feedback_id:
            sheet.delete_rows(row_index, 1)
            workbook.save(path)
            return True
    return False


def feedback_table_path(payload: dict[str, Any]) -> Path:
    feedback_type = str(payload.get("type", "")).strip().lower()
    if feedback_type == "handoff":
        return HANDOFF_FEEDBACK_TABLE_PATH
    if feedback_type == "dislike":
        return UNHELPFUL_FEEDBACK_TABLE_PATH
    return HELPFUL_FEEDBACK_TABLE_PATH


def load_feedback_workbook(path: Path) -> Workbook:
    path.parent.mkdir(parents=True, exist_ok=True)
    if path.exists():
        try:
            workbook = load_workbook(path)
            sheet = workbook.active
            header = [sheet.cell(row=1, column=index + 1).value for index in range(len(FEEDBACK_FIELDS))]
            if header == FEEDBACK_FIELDS:
                hide_feedback_internal_columns(sheet)
                return workbook
            old_header = [sheet.cell(row=1, column=index + 1).value for index in range(len(VISIBLE_FEEDBACK_FIELDS))]
            if old_header == VISIBLE_FEEDBACK_FIELDS:
                sheet.cell(row=1, column=len(FEEDBACK_FIELDS), value="记录ID")
                hide_feedback_internal_columns(sheet)
                return workbook
        except Exception:
            pass
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "反馈表"
    sheet.append(FEEDBACK_FIELDS)
    sheet.column_dimensions["A"].width = 14
    sheet.column_dimensions["B"].width = 22
    sheet.column_dimensions["C"].width = 42
    hide_feedback_internal_columns(sheet)
    workbook.save(path)
    return workbook


def hide_feedback_internal_columns(sheet: Any) -> None:
    sheet.column_dimensions["D"].hidden = True


def natural_sort_key(value: str) -> list[Any]:
    return [int(part) if part.isdigit() else part.lower() for part in re.split(r"(\d+)", value)]


def screenshot_guide_id(rel_dir: Path) -> str:
    safe = re.sub(r"[^0-9A-Za-z\u4e00-\u9fff]+", "-", "-".join(rel_dir.parts)).strip("-")
    return f"screenshot-{safe or 'guide'}"


def normalize_ocr_step_text(value: str) -> str:
    text = re.sub(r"\s+", " ", value).strip()
    text = re.sub(r"(?<=[\u4e00-\u9fff])\s+(?=[\u4e00-\u9fff])", "", text)
    replacements = {
        "占击": "点击",
        "单击": "点击",
        "单兀": "单元",
        "创也": "创建",
        "教師": "教师",
        "教學": "教学",
        "活讥": "活动",
    }
    for source, target in replacements.items():
        text = text.replace(source, target)
    return text


def read_screenshot_sidecar_text(image: Path) -> str:
    for suffix in [".ocr.txt", ".txt"]:
        sidecar = image.with_suffix(suffix)
        if not sidecar.exists():
            continue
        try:
            return normalize_ocr_step_text(sidecar.read_text(encoding="utf-8"))
        except OSError:
            return ""
    return ""


def extract_action_sentence(text: str) -> str:
    for keyword in ["点击", "选择", "进入", "打开", "填写", "上传", "保存", "发布", "确认"]:
        index = text.find(keyword)
        if index < 0:
            continue
        sentence = text[index : index + 48]
        sentence = re.split(
            r"回顶部|关于我们|关于文华|帮助中心|友情链接|国家智慧|联系我|跳转",
            sentence,
            maxsplit=1,
        )[0]
        sentence = sentence[:32].strip(" ，,。.")
        if sentence:
            return f"{sentence}。"
    return ""


def infer_structured_step_text(step_index: int, title_parts: list[str]) -> str:
    leaf_title = title_parts[-1] if title_parts else ""
    parent_title = title_parts[-2] if len(title_parts) >= 2 else ""
    root_title = title_parts[0] if title_parts else ""

    guide_steps = {
        "如何创建AI助手": [
            "进入课程的【AI工作台】，点击助手区域的加号。",
            "进入【管理我的AI助手】页面。",
            "点击【创建AI助手】并选择【手动创建】。",
            "填写助手名称、简介和指令。",
            "按需要配置知识库、工作流和头像，然后保存发布。",
            "返回 AI 助手管理页面，确认新助手已经出现在列表中。",
        ],
        "如何添加AI助手至AI工作台首页": [
            "进入目标课程的【AI工作台】，打开【设置】。",
            "在 AI 助手区域点击【添加AI助手】。",
            "选择要添加到首页的 AI 助手并确认。",
            "返回设置页面，确认该助手状态为【已启用】。",
            "返回 AI 工作台首页，确认刚添加的助手已经显示。",
        ],
        "加入学生-学生扫码": [
            "进入班级管理页面，打开目标班级的学生加入入口。",
            "点击班级二维码，让学生使用优学院 App 扫码加入班级。",
        ],
        "新增AI助手中-选择AI助手": [
            "进入课程的【AI工作台】，点击助手区域的加号。",
            "选择【选择AI助手】，在弹窗中勾选已有助手。",
            "点击【确定】，返回工作台确认助手已经显示。",
        ],
    }
    matched_steps = guide_steps.get(leaf_title)
    if matched_steps and step_index <= len(matched_steps):
        return matched_steps[step_index - 1]

    if root_title == "AI工作台的功能" and leaf_title:
        specific_steps = {
            "公式识别": [
                "在【AI工作台】中找到并点击【公式识别】。",
                "上传包含公式的图片，支持 png、jpg、jpeg 格式。",
                "查看识别结果，可复制公式或重新上传图片。",
            ],
            "备课": [
                "在【AI工作台】中找到并点击【备课】。",
                "按页面提示输入备课要求或上传资料。",
                "查看生成内容，并按需继续编辑或使用。",
            ],
            "学情分析": [
                "在【AI工作台】中找到并点击【学情分析】。",
                "选择需要分析的课程、班级或数据范围。",
                "查看学情分析结果，并按需继续追踪学生情况。",
            ],
            "学生画像": [
                "在【AI工作台】中找到并点击【学生画像】。",
                "选择要查看的学生或班级对象。",
                "查看学生画像内容，辅助后续教学调整。",
            ],
        }
        steps = specific_steps.get(leaf_title)
        if steps and step_index <= len(steps):
            return steps[step_index - 1]
        generic_steps = [
            f"在【AI工作台】中找到并点击【{leaf_title}】。",
            "按页面提示上传资料、输入要求或选择示例。",
            "查看生成或分析结果，并按需继续编辑、复制或使用。",
        ]
        return generic_steps[min(step_index - 1, len(generic_steps) - 1)]

    if parent_title in {"添加教学活动", "添加教学资源"} and leaf_title:
        category = "教学活动" if parent_title == "添加教学活动" else "教学资源"
        if step_index == 1:
            return "进入单元页面，点击【+】添加入口。"
        return f"在弹窗中选择【{category}】里的【{leaf_title}】。"

    return ""


def infer_screenshot_step_text(image: Path, step_index: int, title_parts: list[str]) -> str:
    sidecar_text = read_screenshot_sidecar_text(image)
    leaf_title = title_parts[-1] if title_parts else ""
    parent_title = title_parts[-2] if len(title_parts) >= 2 else ""
    structured_text = infer_structured_step_text(step_index, title_parts)
    if structured_text:
        return structured_text
    if step_index > 1 and parent_title in {"添加教学活动", "添加教学资源"} and leaf_title:
        category = "教学活动" if parent_title == "添加教学活动" else "教学资源"
        return f"在弹窗中选择【{category}】里的【{leaf_title}】。"
    if sidecar_text:
        if leaf_title and leaf_title in sidecar_text and "教学活动" in sidecar_text:
            return f"在弹窗中选择【教学活动】里的【{leaf_title}】。"
        if leaf_title and leaf_title in sidecar_text and "教学资源" in sidecar_text:
            return f"在弹窗中选择【教学资源】里的【{leaf_title}】。"
        if ("加号" in sidecar_text or "+" in sidecar_text) and "教学活动" in sidecar_text:
            return "进入单元页面，点击【+】添加入口。"
        action_sentence = extract_action_sentence(sidecar_text)
        if action_sentence:
            return action_sentence

    stem = image.stem
    text = re.sub(r"(?i)^step\s*\d*", "", stem)
    text = re.sub(r"^\d+[_\-.\s]*", "", text)
    text = text.replace("_", " ").replace("-", " ").strip()
    if text and len(text) <= 24 and not re.fullmatch(r"[0-9a-fA-F]{12,}", text):
        return text
    return f"查看第 {step_index} 张截图，按图中标注完成操作。"


def expand_screenshot_aliases(title_parts: list[str], question: str) -> list[str]:
    base_terms = [title_parts[-1], " - ".join(title_parts), " ".join(title_parts), question, *title_parts]
    action_groups = [
        ("添加", "布置", "发布", "创建", "新建"),
        ("查看", "查询", "看"),
        ("设置", "配置"),
        ("导入", "上传", "添加"),
        ("删除", "移除"),
    ]
    aliases: list[str] = []
    for term in base_terms:
        if term:
            aliases.append(term)
        for group in action_groups:
            matched = next((word for word in group if term.startswith(word)), "")
            if not matched:
                continue
            suffix = term[len(matched):]
            aliases.extend(f"{word}{suffix}" for word in group if suffix)
            aliases.extend(f"如何{word}{suffix}？" for word in group if suffix)
    return list(dict.fromkeys(aliases))


def build_screenshot_guides() -> list[dict[str, Any]]:
    if not SCREENSHOT_DIR.exists():
        return []

    tutorial_dirs = sorted(
        {
            image.parent
            for image in SCREENSHOT_DIR.rglob("*")
            if image.is_file() and image.suffix.lower() in IMAGE_EXTENSIONS
            and "2026用户手册" not in image.relative_to(SCREENSHOT_DIR).parts
        },
        key=lambda path: [natural_sort_key(part) for part in path.relative_to(SCREENSHOT_DIR).parts],
    )

    guides: list[dict[str, Any]] = []
    english_translations: dict[str, Any] = {}
    if SCREENSHOT_EN_TRANSLATIONS_PATH.exists():
        try:
            loaded_translations = json.loads(
                SCREENSHOT_EN_TRANSLATIONS_PATH.read_text(encoding="utf-8")
            )
            if isinstance(loaded_translations, dict):
                english_translations = loaded_translations
        except (OSError, json.JSONDecodeError):
            english_translations = {}
    for tutorial_dir in tutorial_dirs:
        images = sorted(
            [
                image
                for image in tutorial_dir.iterdir()
                if image.is_file() and image.suffix.lower() in IMAGE_EXTENSIONS
            ],
            key=lambda image: natural_sort_key(image.name),
        )
        if not images:
            continue

        rel_dir = tutorial_dir.relative_to(SCREENSHOT_DIR)
        title_parts = list(rel_dir.parts)
        leaf_title = title_parts[-1]
        title = " - ".join(title_parts)
        question = leaf_title if leaf_title.startswith("如何") else f"如何{leaf_title}？"
        aliases = expand_screenshot_aliases(title_parts, question)
        guide_id = screenshot_guide_id(rel_dir)
        english_guide = english_translations.get(guide_id) or {}
        english_steps = english_guide.get("steps") or []
        guides.append(
            {
                "id": guide_id,
                "title": title,
                "titleEn": english_guide.get("title"),
                "question": question,
                "aliases": aliases,
                "path": " > ".join(title_parts),
                "steps": [
                    {
                        "label": f"Step {index}",
                        "text": infer_screenshot_step_text(image, index, title_parts),
                        "textEn": (
                            english_steps[index - 1]
                            if index - 1 < len(english_steps)
                            else ""
                        ),
                        "src": image.relative_to(ROOT).as_posix(),
                    }
                    for index, image in enumerate(images, start=1)
                ],
            }
        )
    try:
        manual_guides = json.loads(MANUAL_2026_GUIDES_PATH.read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError):
        manual_guides = []
    if isinstance(manual_guides, list):
        known_ids = {str(guide.get("id") or "") for guide in guides}
        for guide in manual_guides:
            if not isinstance(guide, dict) or not guide.get("steps"):
                continue
            guide_id = str(guide.get("id") or "")
            if guide_id and guide_id not in known_ids:
                guides.append(guide)
                known_ids.add(guide_id)
    return guides


VIDEO_QUESTION_OVERRIDES = {
    "创建课件": "如何创建课件？",
    "添加章节结构": "如何添加章节结构？",
    "填充课程内容": "如何填充课程内容？",
    "课件设置及发布": "如何进行课件设置及发布？",
    "关联课件": "如何为班课关联教学课件？",
    "添加教学团队": "如何设置教学团队？",
    "发布课程公告": "如何发布公告？",
    "设置学习计划": "如何设置学习计划？",
    "邀请学生加班": "如何邀请学生加入班课？",
    "查看进度成绩": "如何查看进度成绩？",
    "添加资源": "如何添加资源？",
    "个人作业": "如何布置个人作业？",
    "小组作业": "如何布置小组作业？",
    "发布测验": "如何发布测验？",
    "批阅作业": "如何批阅作业？",
    "发布讨论": "如何发布讨论？",
    "试题库-创建试题": "如何添加题目到试题库？",
    "试卷库-添加试卷": "如何添加试卷？",
    "设置课程考核规则": "如何设置课程考核规则？",
    "查看课程分析": "如何查看课程分析？",
    "发起投屏": "如何发起投屏？",
    "结束投屏和导出数据": "如何结束投屏和导出数据？",
    "客服": "如何联系在线客服？",
    "修改个人资料、密码": "如何修改个人资料或密码？",
    "发布考试": "如何安排考试？",
    "考试管理考试分析等": "如何查看考试分析？",
    "课程证书": "如何设置课程证书？",
    "PC端发起直播": "如何在PC端发起直播？",
    "客户端发起直播": "如何在客户端发起直播？",
}


def clean_video_title(path: Path) -> str:
    title = re.sub(r"^\d+[.、_\-\s]*", "", path.stem).strip()
    title = re.sub(r"(?i)pc端?$", "", title).strip()
    return title or path.stem


def build_video_guides() -> list[dict[str, Any]]:
    guides: list[dict[str, Any]] = []
    seen_files: set[str] = set()
    for directory in VIDEO_DIRS:
        if not directory.exists():
            continue
        for video in sorted(
            (
                path
                for path in directory.rglob("*")
                if path.is_file() and path.suffix.lower() in VIDEO_EXTENSIONS
            ),
            key=lambda path: [natural_sort_key(part) for part in path.relative_to(directory).parts],
        ):
            file_key = f"{video.name.lower()}:{video.stat().st_size}"
            if file_key in seen_files:
                continue
            seen_files.add(file_key)
            title = clean_video_title(video)
            question = VIDEO_QUESTION_OVERRIDES.get(title, f"如何{title}？")
            aliases = expand_screenshot_aliases([title], question)
            guides.append(
                {
                    "id": f"video-{len(guides) + 1:03d}",
                    "title": title,
                    "question": question,
                    "aliases": aliases,
                    "src": video.relative_to(ROOT).as_posix(),
                }
            )
    return guides


class AgentHandler(SimpleHTTPRequestHandler):
    def end_headers(self) -> None:
        self.send_header("Cache-Control", "no-store, no-cache, must-revalidate, max-age=0")
        self.send_header("Pragma", "no-cache")
        self.send_header("Expires", "0")
        super().end_headers()

    def _send_json(self, status: int, payload: dict[str, Any]) -> None:
        body = json.dumps(payload, ensure_ascii=False).encode("utf-8")
        self.send_response(status)
        self.send_header("Content-Type", "application/json; charset=utf-8")
        self.send_header("Content-Length", str(len(body)))
        self.end_headers()
        self.wfile.write(body)

    def do_GET(self) -> None:
        request_path = self.path.split("?", 1)[0]
        if request_path == "/api/health":
            configured = bool(
                os.getenv("DEEPSEEK_API_KEY")
                or os.getenv("ULEARNING_TEACHER_ASSISTANT_API_KEY")
                or os.getenv("OPENAI_API_KEY")
            )
            self._send_json(
                200,
                {
                    "ok": True,
                    "configured": configured,
                    "faqCount": len(load_faq_items()),
                    "builtinGuideCount": 0,
                    "screenshotGuideCount": len(build_screenshot_guides()),
                    "videoGuideCount": len(build_video_guides()),
                    "helpKnowledgeCount": 0,
                    "chunkKnowledgeCount": 0,
                    "screenshotKnowledgeCount": 0,
                    "preconditionKnowledgeCount": 0,
                    "knowledgeMode": "deepseek_faq_text_only",
                    "model": os.getenv("DEEPSEEK_MODEL", "deepseek-chat"),
                },
            )
            return
        if request_path == "/api/screenshot-guides":
            self._send_json(200, {"ok": True, "guides": build_screenshot_guides()})
            return
        if request_path == "/api/video-guides":
            self._send_json(200, {"ok": True, "guides": build_video_guides()})
            return
        super().do_GET()

    def do_POST(self) -> None:
        if self.path not in {"/api/chat", "/api/feedback"}:
            self._send_json(404, {"error": "Not found"})
            return

        try:
            length = int(self.headers.get("Content-Length", "0"))
            body = self.rfile.read(length).decode("utf-8")
            payload = json.loads(body) if body else {}
            if self.path == "/api/feedback":
                target_path = feedback_table_path(payload)
                if str(payload.get("action", "")).strip().lower() == "delete":
                    deleted = delete_feedback(payload)
                    self._send_json(200, {"ok": True, "deleted": deleted, "file": target_path.name})
                else:
                    append_feedback(payload)
                    self._send_json(200, {"ok": True, "file": target_path.name})
                return

            message = str(payload.get("message", "")).strip()
            if not message:
                self._send_json(400, {"error": "message is required"})
                return
            language = str(payload.get("language", "zh")).strip().lower()
            self._send_json(200, answer_question_localized(message, language))
        except Exception as exc:
            self._send_json(500, {"error": str(exc)})


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--port", type=int, default=None)
    args = parser.parse_args()

    load_env_file(ENV_PATH)
    port = args.port or int(os.getenv("PORT", "8012"))
    server = ThreadingHTTPServer(("127.0.0.1", port), AgentHandler)
    print(f"Serving on http://127.0.0.1:{port}")
    server.serve_forever()


if __name__ == "__main__":
    main()
