from __future__ import annotations

import json
import re
from pathlib import Path

from PIL import Image
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).parent
SRC = ROOT / "build_cache" / "ulearning_docs_source"
PAGES = SRC / "pages"
PUBLIC = SRC / "public"
IMAGE_CACHE = ROOT / "build_cache" / "help_kb_images"
OUTPUT = ROOT / "help知识库.docx"

PRODUCTS = [
    ("ulearning", "学习管理系统（ULearning）"),
    ("utest", "考试系统（UTest）"),
    ("uclass", "智慧教室（UClass）"),
    ("ulcms", "录播巡课和资源点播平台（ULCMS）"),
]


def set_run_font(run, name="Microsoft YaHei", size=11, bold=None, color=None):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def set_cell_shading(paragraph, fill):
    p_pr = paragraph._p.get_or_add_pPr()
    shd = p_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        p_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run()
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld_sep = OxmlElement("w:fldChar")
    fld_sep.set(qn("w:fldCharType"), "separate")
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_begin, instr, fld_sep, fld_end])
    set_run_font(run, size=9, color="777777")


def configure_document(doc: Document):
    sec = doc.sections[0]
    sec.page_width = Inches(8.5)
    sec.page_height = Inches(11)
    sec.top_margin = Inches(0.75)
    sec.bottom_margin = Inches(0.75)
    sec.left_margin = Inches(0.8)
    sec.right_margin = Inches(0.8)
    sec.header_distance = Inches(0.35)
    sec.footer_distance = Inches(0.35)

    normal = doc.styles["Normal"]
    normal.font.name = "Microsoft YaHei"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.2

    for name, size, before, after, color in [
        ("Heading 1", 16, 16, 8, "2E74B5"),
        ("Heading 2", 13, 12, 6, "2E74B5"),
        ("Heading 3", 11.5, 8, 4, "1F4D78"),
        ("Heading 4", 10.5, 6, 3, "1F4D78"),
    ]:
        style = doc.styles[name]
        style.font.name = "Microsoft YaHei"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    for style_name in ["List Bullet", "List Number"]:
        style = doc.styles[style_name]
        style.font.name = "Microsoft YaHei"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(10.5)
        style.paragraph_format.left_indent = Inches(0.38)
        style.paragraph_format.first_line_indent = Inches(-0.19)
        style.paragraph_format.space_after = Pt(3)
        style.paragraph_format.line_spacing = 1.2

    if "KB Source" not in [s.name for s in doc.styles]:
        src_style = doc.styles.add_style("KB Source", WD_STYLE_TYPE.PARAGRAPH)
        src_style.font.name = "Microsoft YaHei"
        src_style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        src_style.font.size = Pt(8)
        src_style.font.color.rgb = RGBColor(0x66, 0x66, 0x66)
        src_style.paragraph_format.space_after = Pt(7)

    if "KB Callout" not in [s.name for s in doc.styles]:
        style = doc.styles.add_style("KB Callout", WD_STYLE_TYPE.PARAGRAPH)
        style.font.name = "Microsoft YaHei"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(10)
        style.font.color.rgb = RGBColor(0x33, 0x33, 0x33)
        style.paragraph_format.left_indent = Inches(0.15)
        style.paragraph_format.right_indent = Inches(0.15)
        style.paragraph_format.space_before = Pt(4)
        style.paragraph_format.space_after = Pt(6)

    if "KB Code" not in [s.name for s in doc.styles]:
        style = doc.styles.add_style("KB Code", WD_STYLE_TYPE.PARAGRAPH)
        style.font.name = "Consolas"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(9)
        style.paragraph_format.left_indent = Inches(0.15)
        style.paragraph_format.space_after = Pt(2)

    for section in doc.sections:
        add_page_number(section.footer.paragraphs[0])


def clean_inline(text: str) -> str:
    text = re.sub(r"\[([^\]]+)\]\(([^)]+)\)", r"\1（\2）", text)
    text = re.sub(r"[*_]{1,3}([^*_]+)[*_]{1,3}", r"\1", text)
    text = re.sub(r"`([^`]+)`", r"\1", text)
    text = text.replace("<br/>", " ").replace("<br>", " ")
    text = re.sub(r"<[^>]+>", "", text)
    return re.sub(r"\s+", " ", text).strip()


def title_from_mdx(text: str, path: Path) -> str:
    for line in text.splitlines():
        m = re.match(r"^\s*#\s+(.+?)\s*$", line)
        if m:
            return clean_inline(m.group(1))
    m = re.search(r"(?m)^title:\s*(.+?)\s*$", text)
    if m:
        return m.group(1).strip().strip("'\"")
    return path.stem.replace(".zh-CN", "")


def route_for(path: Path) -> str:
    rel = path.relative_to(PAGES).as_posix()
    rel = re.sub(r"\.zh-CN\.mdx$", "", rel)
    return "https://help.ulearning.app/" + rel


def resolve_image(url: str) -> Path | None:
    url = url.strip().split("?")[0].split("#")[0]
    if url.startswith("http://") or url.startswith("https://"):
        return None
    candidate = PUBLIC / url.lstrip("/")
    if candidate.exists():
        return candidate
    known_source_typos = {
        "/Reviewers_pics/revpic3.png": "/Reviewers_pics/zh/revpic3.png",
        "/img/ulearning/teacher/exam011.png": "/img/ulearning/teacher/exam11.png",
    }
    corrected = known_source_typos.get(url)
    if corrected:
        candidate = PUBLIC / corrected.lstrip("/")
        if candidate.exists():
            return candidate
    return None


def compressed_image(path: Path) -> Path | None:
    IMAGE_CACHE.mkdir(exist_ok=True)
    key = re.sub(r"[^A-Za-z0-9_.-]", "_", str(path.relative_to(PUBLIC)))
    out = IMAGE_CACHE / (Path(key).stem + "_" + str(abs(hash(str(path))))[-8:] + ".jpg")
    if out.exists():
        return out
    try:
        with Image.open(path) as img:
            img = img.convert("RGB")
            max_width = 1400
            if img.width > max_width:
                height = max(1, round(img.height * max_width / img.width))
                img = img.resize((max_width, height), Image.Resampling.LANCZOS)
            img.save(out, "JPEG", quality=72, optimize=True, progressive=True)
        return out
    except Exception:
        return None


def add_image(doc: Document, url: str, alt: str, missing: list[str]):
    src = resolve_image(url)
    if not src:
        p = doc.add_paragraph(style="KB Source")
        p.add_run(f"[图片未嵌入] {alt or url}：{url}")
        missing.append(url)
        return
    img = compressed_image(src)
    if not img:
        missing.append(url)
        return
    try:
        with Image.open(img) as im:
            ratio = im.height / max(im.width, 1)
        width = 6.35
        if ratio > 1.7:
            width = 4.3
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(3)
        p.paragraph_format.space_after = Pt(6)
        run = p.add_run()
        run.add_picture(str(img), width=Inches(width))
        if alt:
            cap = doc.add_paragraph(style="KB Source")
            cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
            cap.add_run(alt)
    except Exception:
        missing.append(url)


def parse_article(doc: Document, path: Path, missing: list[str]):
    raw = path.read_text(encoding="utf-8-sig")
    raw = re.sub(r"^---\s*\n.*?\n---\s*\n", "", raw, flags=re.S)
    raw = re.sub(r"(?m)^import\s+.*?;\s*$", "", raw)
    title = title_from_mdx(raw, path)

    doc.add_heading(title, level=1)
    psrc = doc.add_paragraph(style="KB Source")
    psrc.add_run("来源：" + route_for(path))

    in_code = False
    in_callout = False
    first_h1_skipped = False
    paragraph_buffer: list[str] = []

    def flush():
        if paragraph_buffer:
            text = clean_inline(" ".join(paragraph_buffer))
            paragraph_buffer.clear()
            if text:
                p = doc.add_paragraph(style="KB Callout" if in_callout else None)
                if in_callout:
                    set_cell_shading(p, "F4F6F9")
                    r = p.add_run("提示：")
                    set_run_font(r, bold=True, size=10, color="1F4D78")
                p.add_run(text)

    for original in raw.splitlines():
        line = original.rstrip()
        stripped = line.strip()
        if stripped.startswith("```"):
            flush()
            in_code = not in_code
            continue
        if in_code:
            p = doc.add_paragraph(style="KB Code")
            p.add_run(line)
            set_cell_shading(p, "F2F4F7")
            continue
        if "<Callout" in stripped:
            flush()
            in_callout = True
            continue
        if "</Callout>" in stripped:
            flush()
            in_callout = False
            continue
        if not stripped or stripped == "---":
            flush()
            continue

        image_matches = list(re.finditer(r"!\[([^\]]*)\]\(([^)]+)\)", stripped))
        if image_matches:
            before = re.sub(r"!\[[^\]]*\]\([^)]+\)", "", stripped).strip()
            if before:
                paragraph_buffer.append(before)
            flush()
            for m in image_matches:
                add_image(doc, m.group(2), clean_inline(m.group(1)), missing)
            continue

        hm = re.match(r"^(#{1,6})\s+(.+)$", stripped)
        if hm:
            flush()
            level = len(hm.group(1))
            heading_text = clean_inline(hm.group(2))
            if level == 1 and not first_h1_skipped and heading_text == title:
                first_h1_skipped = True
                continue
            doc.add_heading(heading_text, level=min(max(level, 2), 4))
            continue

        bm = re.match(r"^[-*+]\s+(.+)$", stripped)
        if bm:
            flush()
            p = doc.add_paragraph(style="List Bullet")
            p.add_run(clean_inline(bm.group(1)))
            continue
        nm = re.match(r"^\d+[.)、]\s*(.+)$", stripped)
        if nm:
            flush()
            p = doc.add_paragraph(style="List Number")
            p.add_run(clean_inline(nm.group(1)))
            continue
        if stripped.startswith(">"):
            flush()
            p = doc.add_paragraph(style="KB Callout")
            set_cell_shading(p, "F4F6F9")
            p.add_run(clean_inline(stripped.lstrip("> ")))
            continue
        if re.fullmatch(r"<[^>]+>", stripped):
            continue
        paragraph_buffer.append(stripped)
    flush()


def article_files(product: str) -> list[Path]:
    return sorted((PAGES / product).rglob("*.zh-CN.mdx"), key=lambda p: p.as_posix().lower())


def main():
    doc = Document()
    configure_document(doc)

    title_p = doc.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_p.paragraph_format.space_before = Pt(110)
    title_p.paragraph_format.space_after = Pt(18)
    r = title_p.add_run("help知识库")
    set_run_font(r, size=28, bold=True, color="1F4D78")

    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = sub.add_run("优学院帮助中心 · 原始知识点全量收录版")
    set_run_font(r, size=14, color="555555")

    note = doc.add_paragraph()
    note.alignment = WD_ALIGN_PARAGRAPH.CENTER
    note.paragraph_format.space_before = Pt(30)
    r = note.add_run("当前阶段：知识库填充（保留原始颗粒度、重复内容与截图，暂不整合）")
    set_run_font(r, size=10.5, color="777777")

    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    meta.paragraph_format.space_before = Pt(12)
    r = meta.add_run("来源：https://help.ulearning.app/  |  抓取日期：2026-06-24")
    set_run_font(r, size=9, color="777777")

    doc.add_page_break()
    doc.add_heading("收录范围", level=1)
    total = 0
    for product, label in PRODUCTS:
        count = len(article_files(product))
        total += count
        p = doc.add_paragraph(style="List Bullet")
        p.add_run(f"{label}：{count} 篇")
    doc.add_paragraph(f"合计：{total} 篇中文产品帮助文档。")
    doc.add_paragraph("说明：本版按网页文章逐篇入库；不合并相似问题，不删除重复步骤，不重构知识体系。")

    missing: list[str] = []
    for product_index, (product, label) in enumerate(PRODUCTS):
        doc.add_page_break()
        h = doc.add_paragraph()
        h.paragraph_format.space_before = Pt(80)
        h.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = h.add_run(label)
        set_run_font(r, size=23, bold=True, color="2E74B5")
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run(f"本部分共 {len(article_files(product))} 篇")
        doc.add_page_break()
        for path in article_files(product):
            parse_article(doc, path, missing)

    doc.core_properties.title = "help知识库"
    doc.core_properties.subject = "优学院帮助中心原始知识点全量收录"
    doc.core_properties.author = "OpenAI Codex"
    doc.core_properties.comments = "生成于知识库填充阶段，未做内容整合。"
    doc.save(OUTPUT)
    report = {
        "output": str(OUTPUT),
        "articles": sum(len(article_files(p)) for p, _ in PRODUCTS),
        "missing_images": sorted(set(missing)),
        "size_bytes": OUTPUT.stat().st_size,
    }
    report_path = ROOT / "reports" / "help知识库_生成报告.json"
    report_path.parent.mkdir(parents=True, exist_ok=True)
    report_path.write_text(
        json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    print(json.dumps(report, ensure_ascii=False))


if __name__ == "__main__":
    main()
